"""
Deterministic synthetic SmPC fixture WITH embedded images (P1-IMG-1..4).

Produces `tests/fixtures/synthetic_smpc_images.docx`: the same fictional
"Synthex 50 mg" SmPC as `gen_synthetic_smpc.py`, plus four inline pictures
placed so that every P1-IMG code path is exercised exactly once:

  #  | section | format | alt text                         | exercises
  ---+---------+--------+----------------------------------+------------------------------
  1  | 2       | PNG    | present                          | IMG-EMBED (web-safe, alt kept)
  2  | 4.2     | PNG    | ABSENT                           | IMG-ALT-MISSING -> alt="Figure 2"
  3  | 5.1     | TIFF   | present                          | IMG-RASTERISE (TIFF -> PNG)
  4  | 5.2     | PNG    | present, SAME BYTES as #1        | IMG-EMBED dedupe -> same Binary id

All pixels are synthetic geometry (no fonts, no text) so the PNG/TIFF bytes
are byte-identical on every run of this generator on every platform.

Usage:
    python3 tests/fixtures/gen_synthetic_smpc_images.py

Refs:
- FEATURE_SPEC.md §5 P1-IMG-1..4; CLAUDE.md §7.2 (golden corpus), §4.4
  (reproducible runs), §10 #15 (fixture content is intentionally fictional).
"""
from __future__ import annotations

import hashlib
import io
import os
import sys

from docx.shared import Mm
from PIL import Image, ImageDraw

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_synthetic_smpc import build_document  # noqa: E402  (same directory)


# --------------------------------------------------------------------------
# Deterministic synthetic images
# --------------------------------------------------------------------------

def _png_structure() -> bytes:
    """120x60 PNG: a fictional 'chemical structure' made of hexagons/lines."""
    im = Image.new("RGB", (120, 60), "white")
    d = ImageDraw.Draw(im)
    hexagon = [(20, 30), (30, 13), (50, 13), (60, 30), (50, 47), (30, 47)]
    d.polygon(hexagon, outline="black")
    d.line([(60, 30), (80, 30)], fill="black", width=2)
    d.ellipse([(80, 20), (100, 40)], outline="black", width=2)
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=False, compress_level=6)
    return buf.getvalue()


def _png_flowchart() -> bytes:
    """160x80 PNG: two boxes and an arrow (a fictional dosing flow)."""
    im = Image.new("RGB", (160, 80), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([(10, 20), (60, 60)], outline="black", width=2)
    d.rectangle([(100, 20), (150, 60)], outline="black", width=2)
    d.line([(60, 40), (100, 40)], fill="black", width=2)
    d.polygon([(100, 40), (92, 35), (92, 45)], fill="black")
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=False, compress_level=6)
    return buf.getvalue()


def _tiff_chart() -> bytes:
    """140x70 uncompressed TIFF: bars (a fictional PK chart). Non-web-safe MIME."""
    im = Image.new("RGB", (140, 70), "white")
    d = ImageDraw.Draw(im)
    for i, h in enumerate((20, 35, 50, 42, 28)):
        x0 = 10 + i * 25
        d.rectangle([(x0, 65 - h), (x0 + 15, 65)], fill="black")
    d.line([(5, 65), (135, 65)], fill="black", width=1)
    buf = io.BytesIO()
    im.save(buf, format="TIFF", compression=None)
    return buf.getvalue()


IMAGES = {
    # key: (bytes, filename, alt_text or None, section heading to attach to)
    "fig1": (_png_structure(), "fig1.png",
             "Chemical structure of synthexine (fictional)", "2."),
    "fig2": (_png_flowchart(), "fig2.png",
             None, "4.2"),
    "fig3": (_tiff_chart(), "fig3.tiff",
             "Mean plasma concentration over time (fictional)", "5.1"),
    "fig4": (_png_structure(), "fig4.png",   # SAME BYTES as fig1 -> dedupe
             "Chemical structure of synthexine, repeated (fictional)", "5.2"),
}


# --------------------------------------------------------------------------
# DOCX assembly
# --------------------------------------------------------------------------

def _find_heading_paragraph(doc, prefix: str):
    """Return the first bold paragraph whose text starts with `prefix`."""
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith(prefix) and p.runs and p.runs[0].bold:
            return p
    raise KeyError(f"heading starting with {prefix!r} not found")


def _next_heading_after(doc, para):
    """Return the first bold heading paragraph after `para` (or None)."""
    seen = False
    for p in doc.paragraphs:
        if p._p is para._p:
            seen = True
            continue
        if seen and p.runs and p.runs[0].bold and p.text.strip():
            return p
    return None


def _insert_picture_at_section_end(doc, heading_prefix: str, data: bytes,
                                   filename: str, alt: str | None) -> None:
    heading = _find_heading_paragraph(doc, heading_prefix)
    nxt = _next_heading_after(doc, heading)
    if nxt is None:
        target = doc.add_paragraph()
    else:
        target = nxt.insert_paragraph_before()
    run = target.add_run()
    shape = run.add_picture(io.BytesIO(data), width=Mm(40))
    # python-docx exposes docPr; `descr` is what Word shows as "Alt Text" and
    # what mammoth surfaces as <img alt="...">.
    doc_pr = shape._inline.docPr
    doc_pr.set("name", filename)
    if alt is not None:
        doc_pr.set("descr", alt)
    else:
        # Make sure no alt text sneaks in from the template.
        if "descr" in doc_pr.attrib:
            del doc_pr.attrib["descr"]


def build_document_with_images():
    doc = build_document()
    for key in ("fig1", "fig2", "fig3", "fig4"):
        data, filename, alt, heading = IMAGES[key]
        _insert_picture_at_section_end(doc, heading, data, filename, alt)
    return doc


def main() -> int:
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "synthetic_smpc_images.docx")
    doc = build_document_with_images()
    doc.save(out_path)
    print(f"wrote {out_path} ({os.path.getsize(out_path):,} bytes)")
    for key, (data, filename, alt, heading) in IMAGES.items():
        print(f"  {key}: {filename:9s} section {heading:4s} "
              f"sha256={hashlib.sha256(data).hexdigest()[:16]} "
              f"alt={'yes' if alt else 'NO'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
