"""
Deterministic synthetic SmPC fixture WITH a table (P0-8a).

Produces `tests/fixtures/synthetic_smpc_table.docx`: the same fictional
"Synthex 50 mg" SmPC as `gen_synthetic_smpc.py`, plus a 3x2 dosing table in
section 4.2 so the table-border path (mapper + Phase 1 fixer) is exercised
end-to-end. The baseline fixture `synthetic_smpc.docx` has no table and is
left untouched (golden-corpus discipline, CLAUDE.md §5.8).

Usage:
    python3 tests/fixtures/gen_synthetic_smpc_table.py

Refs: FEATURE_SPEC.md §5 P0-8a; CLAUDE.md §10 #15 (fictional content).
"""
from __future__ import annotations

import os
import sys

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_synthetic_smpc import build_document  # noqa: E402

ROWS = [("Population", "Dose (fictional)"),
        ("Adults", "One 50 mg tablet once daily"),
        ("Elderly", "One 50 mg tablet once daily; no adjustment")]


def build_document_table() -> Document:
    doc = build_document()
    paras = doc.paragraphs
    idx = next(i for i, p in enumerate(paras) if p.text.strip().startswith("4.2 Posology"))
    anchor = paras[idx + 1]
    table = doc.add_table(rows=len(ROWS), cols=2)
    for r, (a, b) in enumerate(ROWS):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    anchor._p.addnext(table._tbl)  # place the table right after the first 4.2 paragraph
    return doc


def main() -> int:
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "synthetic_smpc_table.docx")
    build_document_table().save(out_path)
    print(f"wrote {out_path} ({os.path.getsize(out_path):,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
