"""
Deterministic synthetic SmPC fixture whose QRD section numbers exist ONLY as
Word automatic numbering (P0-3a).

Produces `tests/fixtures/synthetic_smpc_autonum.docx`: the same fictional
"Synthex 50 mg" SmPC as `gen_synthetic_smpc.py`, but every section heading is
stored as plain text ("NAME OF THE MEDICINAL PRODUCT", "Therapeutic
indications", …) carrying a `w:numPr` that references a two-level list
definition (`%1.` / `%1.%2`) — exactly how Word saves an SmPC authored with
its built-in outline numbering. Word paints "1.", "4.1" on screen; the
characters are not in the document text. This is the shape that made a real
MAH SmPC fail the P0-2 gate on 2026-09-09.

Also included, and deliberately NOT a heading: an auto-numbered body list
("Take one tablet…", "Do not crush…") inside section 4.2, using a separate
single-level list definition. The pre-pass must leave it untouched so it still
renders as <ol><li>.

Usage:
    python3 tests/fixtures/gen_synthetic_smpc_autonum.py

Refs:
- FEATURE_SPEC.md §5 P0-3a; CLAUDE.md §7.2 (golden corpus), §10 #15
  (fixture content is intentionally fictional).
"""
from __future__ import annotations

import copy
import os
import re
import sys

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_synthetic_smpc import _SECTIONS  # noqa: E402  (same directory)

# Body list that must survive as a real Word list (not a heading).
_BODY_LIST_SECTION = "4.2 Posology and method of administration"
_BODY_LIST_ITEMS = [
    "Take one tablet with a glass of water (fictional instruction).",
    "Do not crush or chew the tablet (fictional instruction).",
    "If a dose is missed, take the next dose at the usual time (fictional instruction).",
]

_QRD_ABSTRACT = """
<w:abstractNum {ns} w:abstractNumId="{aid}">
  <w:multiLevelType w:val="multilevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/></w:lvl>
  <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/><w:lvlJc w:val="left"/></w:lvl>
  <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/></w:lvl>
</w:abstractNum>
"""

_BODY_ABSTRACT = """
<w:abstractNum {ns} w:abstractNumId="{aid}">
  <w:multiLevelType w:val="singleLevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/></w:lvl>
</w:abstractNum>
"""


def _add_list_definition(doc: Document, abstract_xml: str) -> int:
    """Append an abstractNum + num to numbering.xml; return the new numId."""
    numbering = doc.part.numbering_part.element
    existing_abs = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    aid = max(existing_abs, default=-1) + 1
    nid = max(existing_num, default=0) + 1
    abstract = parse_xml(abstract_xml.format(ns=nsdecls("w"), aid=aid))
    # abstractNum elements must precede num elements.
    nums = numbering.findall(qn("w:num"))
    if nums:
        nums[0].addprevious(abstract)
    else:
        numbering.append(abstract)
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/></w:num>')
    numbering.append(num)
    return nid


def _set_numbering(paragraph, num_id: int, ilvl: int) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    numPr = parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    )
    pPr.append(numPr)


_HEADING_RE = re.compile(r"^(\d+)(?:\.(\d+))?\.?\s+(.*)$")


def build_document_autonum() -> Document:
    doc = Document()
    for p_text in ("ANNEX I", "SUMMARY OF PRODUCT CHARACTERISTICS"):
        p = doc.add_paragraph()
        p.add_run(p_text).bold = True
        doc.add_paragraph()

    qrd_num = _add_list_definition(doc, _QRD_ABSTRACT)
    body_num = _add_list_definition(doc, _BODY_ABSTRACT)

    for heading, body in _SECTIONS:
        m = _HEADING_RE.match(heading)
        assert m, heading
        _major, minor, title = m.groups()
        p_h = doc.add_paragraph()
        p_h.add_run(title).bold = True          # NO number in the text …
        _set_numbering(p_h, qrd_num, 1 if minor else 0)  # … only Word numbering
        for para in body.split("\n"):
            if para.strip():
                doc.add_paragraph(para)
        if heading == _BODY_LIST_SECTION:
            for item in _BODY_LIST_ITEMS:
                p_i = doc.add_paragraph(item)
                _set_numbering(p_i, body_num, 0)
    return doc


def main() -> int:
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "synthetic_smpc_autonum.docx")
    build_document_autonum().save(out_path)
    print(f"wrote {out_path} ({os.path.getsize(out_path):,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
